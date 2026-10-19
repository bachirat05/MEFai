from __future__ import annotations

import unicodedata
from collections.abc import Callable
from pathlib import Path

from docx import Document
from docx.table import Table
from pypdf import PdfReader


def normalize_text(value: str) -> str:
    """Minuscules + suppression des accents (« Congé » -> « conge »)."""
    return (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )


# ------------------------------------------------------------------
# TXT
# ------------------------------------------------------------------


def read_txt(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8-sig", errors="ignore")


# ------------------------------------------------------------------
# PDF
# ------------------------------------------------------------------


def read_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    pages = [
        f"\n--- PAGE {number} ---\n{text}"
        for number, page in enumerate(reader.pages, start=1)
        if (text := page.extract_text() or "").strip()
    ]
    return "\n".join(pages)


# ------------------------------------------------------------------
# DOCX
# ------------------------------------------------------------------


def _clean_cell_text(text: str) -> str:
    return " ".join(line.strip() for line in text.splitlines() if line.strip())


def _row_values(row) -> list[str]:
    """Texte de chaque cellule d'une ligne.

    Une cellule fusionnée apparaît plusieurs fois dans `row.cells` : on ne la
    garde qu'à sa première occurrence et on met "" ensuite. Les colonnes restent
    ainsi alignées avec les en-têtes (contrairement à un dédoublonnage par valeur,
    qui supprimerait aussi deux cellules distinctes contenant le même texte).
    """
    seen: set[int] = set()
    values: list[str] = []
    for cell in row.cells:
        key = id(cell._tc)
        values.append("" if key in seen else _clean_cell_text(cell.text))
        seen.add(key)
    return values


def _table_to_text(table: Table) -> list[str]:
    rows = [r for r in map(_row_values, table.rows) if any(r)]
    if not rows:
        return []

    headers, body = rows[0], rows[1:]
    if not body:
        return [" | ".join(h for h in headers if h)]

    blocks: list[str] = []
    for row in body:
        pairs = []
        for index, value in enumerate(row):
            if not value:
                continue
            header = headers[index] if index < len(headers) and headers[index] else f"Colonne {index + 1}"
            if value != header:
                pairs.append(f"{header}: {value}")
        if pairs:
            blocks.append("\n".join(pairs))
    return blocks


def read_docx(file_path: Path) -> str:
    document = Document(str(file_path))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        parts.extend(_table_to_text(table))
    return "\n\n".join(parts)


# ------------------------------------------------------------------
# Dispatch / découverte
# ------------------------------------------------------------------

# Source unique de vérité : ajouter un format = ajouter une ligne ici.
_READERS: dict[str, Callable[[Path], str]] = {
    ".txt": read_txt,
    ".pdf": read_pdf,
    ".docx": read_docx,
}

SUPPORTED_EXTENSIONS = frozenset(_READERS)


def read_document(file_path: Path) -> str:
    extension = file_path.suffix.lower()
    reader = _READERS.get(extension)
    if reader is None:
        raise ValueError(f"Type de fichier non supporté : {extension}")
    return reader(file_path)


def list_files(directory: Path) -> list[Path]:
    """Fichiers supportés sous `directory` (récursif, triés).

    Ignore les fichiers cachés et les fichiers de verrouillage Word (`~$...docx`),
    qui ont l'extension .docx mais ne sont pas lisibles.
    """
    return sorted(
        path
        for path in directory.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not path.name.startswith(("~$", "."))
    )