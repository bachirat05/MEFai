from __future__ import annotations

import logging
import math
import tempfile

from pathlib import Path

import httpx
import pytest
from unittest.mock import AsyncMock, MagicMock, patch
from docx import Document

from core.config import BASE_DIR, Settings, get_settings
from schemas.document import DocumentCategory, DocumentIngestion, DocumentMetadata, DocumentSource, DocumentSearchQuery
from services.rag.document_ingestion import build_file_ingestions, detect_category, read_docx
from schemas.workflow import DemandeConge, TypeConge
from services.rag.rag_service import RAGInitializationError, RAGService
from services.workflows.n8n_service import N8nHTTPError, N8nService, N8nServiceError, N8nTimeoutError

logger = logging.getLogger(__name__)

settings = get_settings()


def test_chroma_persist_dir_is_resolved_from_project_root() -> None:
    settings_with_relative_chroma = Settings(
        chroma_persist_dir="./chroma_db",
    )

    assert settings_with_relative_chroma.chroma_persist_dir == str(
        (
            BASE_DIR
            / "chroma_db"
        ).resolve()
    )


def test_prime_transport_filename_detected_as_benefits() -> None:
    category = detect_category(
        Path("test_prime_transport.txt")
    )

    assert category == DocumentCategory.BENEFITS


def test_formation_certification_filename_detected_as_training() -> None:
    category = detect_category(
        Path(
            "Politique_Formation_Certification_Numerique_Demo.docx"
        )
    )

    assert category == DocumentCategory.TRAINING


def test_docx_reader_extracts_paragraphs_and_tables(
    tmp_path: Path,
) -> None:
    file_path = (
        tmp_path
        / "formation_certification.docx"
    )
    document = Document()

    document.add_paragraph(
        "Chaque collaborateur dispose d'un credit indicatif de 4 jours de formation par annee civile."
    )

    table = document.add_table(
        rows=2,
        cols=4,
    )
    table.rows[0].cells[0].text = "Element"
    table.rows[0].cells[1].text = "Regle de demonstration"
    table.rows[0].cells[2].text = "Limite"
    table.rows[0].cells[3].text = "Cellule vide"
    table.rows[1].cells[0].text = "Frais d'examen"
    table.rows[1].cells[1].text = "Remboursement a hauteur de 80 % apres reussite"
    table.rows[1].cells[2].text = "2 500 MAD par certification"
    table.rows[1].cells[3].text = ""

    document.save(
        file_path
    )

    text = read_docx(
        file_path
    )

    assert (
        "4 jours de formation par annee civile"
        in text
    )
    assert "Element: Frais d'examen" in text
    assert (
        "Regle de demonstration: Remboursement a hauteur de 80 % apres reussite"
        in text
    )
    assert (
        "Limite: 2 500 MAD par certification"
        in text
    )
    assert "Cellule vide:" not in text


class MockEmbeddingFunction:
    def __init__(self, dimension: int = 384) -> None:
        self._dimension = dimension

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[float(i + j * 0.001) for i in range(self._dimension)] for j, _ in enumerate(texts)]

    def embed_query(self, text: str) -> list[float]:
        return [float(i * 0.001) for i in range(self._dimension)]


class HashedKeywordEmbeddingFunction:
    def __init__(self, dimension: int = 256) -> None:
        self._dimension = dimension

    def _embed(self, text: str) -> list[float]:
        vector = [0.0] * self._dimension
        normalized = (
            text.lower()
            .replace("é", "e")
            .replace("è", "e")
            .replace("ê", "e")
            .replace("à", "a")
            .replace("ç", "c")
        )
        tokens = [
            token.strip(".,;:!?()[]'\"")
            for token in normalized.split()
        ]
        for token in tokens:
            if len(token) < 3:
                continue
            vector[hash(token) % self._dimension] += 1.0

        norm = math.sqrt(
            sum(value * value for value in vector)
        )
        if norm:
            return [
                value / norm
                for value in vector
            ]

        return vector

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [
            self._embed(text)
            for text in texts
        ]

    def embed_query(self, text: str) -> list[float]:
        return self._embed(text)


@pytest.fixture
def temp_chroma_dir() -> str:
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        yield tmp


@pytest.fixture
def rag_service(temp_chroma_dir: str) -> RAGService:
    embedding = MockEmbeddingFunction(dimension=384)
    svc = RAGService(
        embedding_function=embedding,
        persist_directory=temp_chroma_dir,
        collection_name="test_rh_documents",
    )
    svc.initialize()
    return svc


@pytest.fixture
def keyword_rag_service(temp_chroma_dir: str) -> RAGService:
    svc = RAGService(
        embedding_function=HashedKeywordEmbeddingFunction(),
        persist_directory=temp_chroma_dir,
        collection_name="test_keyword_rh_documents",
    )
    svc.initialize()
    return svc


@pytest.fixture
def populated_rag_service(rag_service: RAGService) -> RAGService:
    doc = DocumentIngestion(
        content="Le congé annuel accorde 25 jours ouvrables par an pour tout salarié à temps plein.",
        metadata=DocumentMetadata(
            category=DocumentCategory.RH_POLICY,
            source=DocumentSource.INTERNAL_WIKI,
            department="RH",
            tags=["congé", "annuel", "droit"],
        ),
    )
    rag_service.ingest(doc)
    return rag_service


@pytest.fixture
def n8n_service() -> N8nService:
    return N8nService(
        webhook_url="http://localhost:5678/webhook/congé-test",
        secret_token="test-secret-token",
        timeout_seconds=settings.n8n_timeout_seconds,
    )


@pytest.mark.asyncio
async def test_rag_retrieval_with_documents(populated_rag_service: RAGService) -> None:
    query = DocumentSearchQuery(
        query="Combien de jours de congé annuel ?",
        category=DocumentCategory.RH_POLICY,
        top_k=5,
        similarity_threshold=0.0,
    )
    results = populated_rag_service.search(query)

    assert results.total_found >= 1, "Aucun document retourné — vérifiez que l'ingestion a été effectuée"
    assert len(results.results) > 0, "La liste de résultats est vide"
    assert results.results[0].score >= 0.0, "Le score de similarité doit être positif ou nul"
    assert len(results.results[0].content) > 0, "Le contenu du document ne doit pas être vide"


def test_dynamic_uploaded_txt_document_is_immediately_retrievable(
    tmp_path: Path,
    keyword_rag_service: RAGService,
) -> None:
    file_path = tmp_path / "mobilite_interne_temporaire.txt"
    file_path.write_text(
        (
            "Politique fictive de mobilite interne.\n"
            "Le delai de traitement d'une demande de mobilite interne "
            "est de 17 jours ouvres.\n"
            "Le code de reference est MOB-TEST-917."
        ),
        encoding="utf-8",
    )

    ingestions = build_file_ingestions(
        file_path
    )
    keyword_rag_service.ingest_batch(
        ingestions
    )

    questions = {
        "Quel est le delai de traitement d'une demande de mobilite interne ?": "17 jours ouvres",
        "Combien de temps prend une demande de mobilite interne ?": "17 jours ouvres",
        "Quel est le code de reference de la politique de mobilite ?": "MOB-TEST-917",
    }

    for question, expected in questions.items():
        response = keyword_rag_service.search(
            DocumentSearchQuery(
                query=question,
                category=None,
                top_k=10,
                similarity_threshold=0.0,
            )
        )

        combined = "\n".join(
            result.content
            for result in response.results
        )

        assert expected in combined


def test_unknown_category_document_is_found_by_global_search(
    tmp_path: Path,
    keyword_rag_service: RAGService,
) -> None:
    file_path = tmp_path / "document_sans_categorie_connue.txt"
    file_path.write_text(
        (
            "Note interne experimentale.\n"
            "La fenetre de reponse du dispositif Alpha est de 42 heures.\n"
            "Le marqueur documentaire est ALPHA-UNKNOWN-42."
        ),
        encoding="utf-8",
    )

    ingestions = build_file_ingestions(
        file_path
    )

    assert ingestions[0].metadata.category == DocumentCategory.OTHER

    keyword_rag_service.ingest_batch(
        ingestions
    )

    response = keyword_rag_service.search(
        DocumentSearchQuery(
            query="Quel est le marqueur documentaire du dispositif Alpha ?",
            category=None,
            top_k=10,
            similarity_threshold=0.0,
        )
    )

    assert any(
        "ALPHA-UNKNOWN-42" in result.content
        for result in response.results
    )


def test_precise_code_and_amount_document_values_are_retrievable(
    tmp_path: Path,
    keyword_rag_service: RAGService,
) -> None:
    file_path = tmp_path / "mobilite_valeurs_precises.txt"
    file_path.write_text(
        (
            "Politique fictive de mobilite interne.\n"
            "Le code de reference est XYZ-DEMO-481.\n"
            "Une indemnite de mobilite de 1 725 MAD est prevue."
        ),
        encoding="utf-8",
    )

    ingestions = build_file_ingestions(
        file_path
    )
    keyword_rag_service.ingest_batch(
        ingestions
    )

    expectations = {
        "Quel est le code de reference ?": "XYZ-DEMO-481",
        "Quel montant est prevu pour l'indemnite de mobilite ?": "1 725 MAD",
    }

    for question, expected in expectations.items():
        response = keyword_rag_service.search(
            DocumentSearchQuery(
                query=question,
                category=None,
                top_k=10,
                similarity_threshold=0.0,
            )
        )
        combined = "\n".join(
            result.content
            for result in response.results
        )

        assert expected in combined


def test_uploaded_hardware_policy_values_are_retrievable_without_category(
    tmp_path: Path,
    keyword_rag_service: RAGService,
) -> None:
    file_path = tmp_path / "politique_materiel_test.txt"
    file_path.write_text(
        (
            "POLITIQUE MATERIEL TEST\n"
            "Renouvellement: tous les 7 ans.\n"
            "Plafond accessoires: 875 MAD.\n"
            "Code: MAT-RAG-731."
        ),
        encoding="utf-8",
    )

    ingestions = build_file_ingestions(
        file_path
    )
    keyword_rag_service.ingest_batch(
        ingestions
    )

    expectations = {
        "Tous les combien d'annees le materiel est-il renouvele ?": "7 ans",
        "Quel est le plafond des accessoires ?": "875 MAD",
        "Quel est le code de reference de la politique materiel ?": "MAT-RAG-731",
    }

    for question, expected in expectations.items():
        response = keyword_rag_service.search(
            DocumentSearchQuery(
                query=question,
                category=None,
                top_k=10,
                similarity_threshold=0.0,
            )
        )
        combined = "\n".join(
            result.content
            for result in response.results
        )

        assert expected in combined


def test_rag_delete_batch_removes_only_target_chunks(rag_service: RAGService) -> None:
    target = DocumentIngestion(
        content="Document temporaire a supprimer du RAG.",
        metadata=DocumentMetadata(
            category=DocumentCategory.RH_POLICY,
            source=DocumentSource.INTERNAL_WIKI,
            department="RH",
            tags=["target"],
        ),
    )
    retained = DocumentIngestion(
        content="Document temporaire qui doit rester dans le RAG.",
        metadata=DocumentMetadata(
            category=DocumentCategory.BENEFITS,
            source=DocumentSource.INTERNAL_WIKI,
            department="RH",
            tags=["retained"],
        ),
    )

    target_id, retained_id = rag_service.ingest_batch(
        [
            target,
            retained,
        ]
    )

    deleted_ids = rag_service.delete_batch(
        [
            target,
        ]
    )

    assert deleted_ids == [target_id]

    existing = rag_service._vectorstore.get(
        ids=[
            target_id,
            retained_id,
        ],
        include=[],
    )

    assert target_id not in existing.get("ids", [])
    assert retained_id in existing.get("ids", [])


def test_rag_delete_by_tag_removes_only_matching_chunks(
    rag_service: RAGService,
) -> None:
    target = DocumentIngestion(
        content="Ancien chunk du document formation certification.",
        metadata=DocumentMetadata(
            category=DocumentCategory.TRAINING,
            source=DocumentSource.INTERNAL_WIKI,
            department="RH",
            tags=[
                "formation_certification",
                ".docx",
            ],
        ),
    )
    retained = DocumentIngestion(
        content="Chunk d'un autre document formation.",
        metadata=DocumentMetadata(
            category=DocumentCategory.TRAINING,
            source=DocumentSource.INTERNAL_WIKI,
            department="RH",
            tags=[
                "autre_document",
                ".docx",
            ],
        ),
    )

    target_id, retained_id = rag_service.ingest_batch(
        [
            target,
            retained,
        ]
    )

    deleted_ids = rag_service.delete_by_tag(
        "formation_certification"
    )

    assert deleted_ids == [target_id]

    existing = rag_service._vectorstore.get(
        ids=[
            target_id,
            retained_id,
        ],
        include=[],
    )

    assert target_id not in existing.get("ids", [])
    assert retained_id in existing.get("ids", [])


@pytest.mark.asyncio
async def test_rag_retrieval_empty_database(temp_chroma_dir: str) -> None:
    embedding = MockEmbeddingFunction(dimension=384)
    svc = RAGService(
        embedding_function=embedding,
        persist_directory=temp_chroma_dir,
        collection_name="test_empty_rh_documents",
    )
    svc.initialize()

    query = DocumentSearchQuery(
        query="Combien de jours de congé ?",
        category=DocumentCategory.RH_POLICY,
        top_k=5,
        similarity_threshold=0.0,
    )
    results = svc.search(query)

    logger.warning("La base ChromaDB est vide — lancez l'ingestion des documents RH avant de faire des recherches.")

    assert results.total_found == 0, "Attendu : aucun résultat car la base est vide"
    assert results.returned_count == 0, "Attendu : aucun résultat retourné"


@pytest.mark.asyncio
async def test_rag_retrieval_non_initialized_service() -> None:
    embedding = MockEmbeddingFunction(dimension=384)
    svc = RAGService(
        embedding_function=embedding,
        persist_directory="/tmp/nonexistent_test_dir",
        collection_name="test_uninitialized",
    )

    query = DocumentSearchQuery(
        query="test",
        top_k=5,
    )

    with pytest.raises(RAGInitializationError):
        svc.search(query)


@pytest.mark.asyncio
async def test_n8n_webhook_success(n8n_service: N8nService) -> None:
    mock_response = AsyncMock()
    mock_response.status_code = 200
    mock_response.headers = {"content-type": "application/json"}

    mock_response.json = MagicMock(
    return_value={
        "statut": "accepted",
        "message": "Demande transmise avec succès au manager",
        "demande_id": "REQ-001",
    }
)

    mock_response.raise_for_status = MagicMock()

    mock_post = AsyncMock(return_value=mock_response)

    with patch("httpx.AsyncClient.post", mock_post):
        demande = DemandeConge(
            employe_id="EMP-001",
            date_debut="2026-08-01",
            date_fin="2026-08-05",
            type_conge=TypeConge.CONGE_ANNUEL,
            motif="Congé annuel pour voyage familial",
        )
        result = await n8n_service.soumettre_demande_conge(
            demande,
            manager_email="manager@example.com",
        )

    assert result is not None, "Le résultat de n8n ne doit pas être None"
    assert result.get("statut") == "accepted", f"Statut attendu 'accepted', obtenu : {result.get('statut')}"
    assert result.get("message") == "Demande transmise avec succès au manager"


@pytest.mark.asyncio
async def test_n8n_webhook_timeout(n8n_service: N8nService) -> None:
    mock_post = AsyncMock(side_effect=httpx.TimeoutException("Connection timed out"))

    with patch("httpx.AsyncClient.post", mock_post):
        demande = DemandeConge(
            employe_id="EMP-002",
            date_debut="2026-08-10",
            date_fin="2026-08-15",
            type_conge=TypeConge.CONGE_MALADIE,
            motif="Congé pour raisons médicales",
        )
        with pytest.raises(N8nTimeoutError) as exc_info:
            await n8n_service.soumettre_demande_conge(
                demande,
                manager_email="manager@example.com",
            )

    assert "timed out" in str(exc_info.value).lower() or "timeout" in str(exc_info.value).lower()


@pytest.mark.asyncio
async def test_n8n_webhook_http_error(n8n_service: N8nService) -> None:
    mock_response = AsyncMock()
    mock_response.status_code = 503
    mock_response.text = "Service Unavailable"

    mock_http_error = httpx.HTTPStatusError(
        "503 Server Error",
        request=MagicMock(),
        response=mock_response,
    )
    mock_post = AsyncMock(side_effect=mock_http_error)

    with patch("httpx.AsyncClient.post", mock_post):
        demande = DemandeConge(
            employe_id="EMP-003",
            date_debut="2026-09-01",
            date_fin="2026-09-10",
            type_conge=TypeConge.CONGE_SANS_SOLDE,
            motif="Congé sans solde personnel",
        )
        with pytest.raises(N8nHTTPError) as exc_info:
            await n8n_service.soumettre_demande_conge(
                demande,
                manager_email="manager@example.com",
            )

    assert exc_info.value.status_code == 503


@pytest.mark.asyncio
async def test_n8n_webhook_connection_error(n8n_service: N8nService) -> None:
    mock_post = AsyncMock(side_effect=httpx.ConnectError("Connection refused"))

    with patch("httpx.AsyncClient.post", mock_post):
        demande = DemandeConge(
            employe_id="EMP-004",
            date_debut="2026-10-01",
            date_fin="2026-10-05",
            type_conge=TypeConge.CONGE_FORMATION,
            motif="Congé pour formation professionnelle",
        )
        with pytest.raises(N8nServiceError):
            await n8n_service.soumettre_demande_conge(
                demande,
                manager_email="manager@example.com",
            )


@pytest.mark.asyncio
async def test_n8n_service_no_webhook_url_configured() -> None:
    service = N8nService(
        webhook_url="",
        secret_token="token",
        timeout_seconds=10.0,
    )

    demande = DemandeConge(
        employe_id="EMP-005",
        date_debut="2026-11-01",
        date_fin="2026-11-05",
        type_conge=TypeConge.AUTRE,
        motif="Autre type de congé",
    )

    with pytest.raises(N8nServiceError, match="n8n_webhook_url is not configured"):
        await service.soumettre_demande_conge(
            demande,
            manager_email="manager@example.com",
        )


@pytest.mark.asyncio
async def test_n8n_invalid_demande_conge() -> None:
    service = N8nService(
        webhook_url="http://localhost:5678/webhook/test",
        secret_token="token",
        timeout_seconds=10.0,
    )

    with pytest.raises(Exception):
        await service.soumettre_demande_conge(
            DemandeConge(
                employe_id="",
                date_debut="2026-12-10",
                date_fin="2026-12-05",
                type_conge=TypeConge.CONGE_ANNUEL,
                motif="",
            ),
            manager_email="manager@example.com",
        )


@pytest.mark.asyncio
async def test_demande_conge_date_fin_inferieure_a_date_debut() -> None:
    with pytest.raises(Exception):
        DemandeConge(
            employe_id="EMP-006",
            date_debut="2026-12-10",
            date_fin="2026-12-05",
            type_conge=TypeConge.CONGE_ANNUEL,
            motif="Test date inversée",

        )

